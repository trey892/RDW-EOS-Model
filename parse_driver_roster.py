"""
Parse McLeod's full "DRIVER LISTING" export (Drive title "Driver Data
<date>.xlsx", SharePoint Sync folder) -- the authoritative Driver Code <->
Name <-> Fleet/Terminal roster this model was missing. Confirmed 2026-08-20:
a genuine daily-ish export (today's pull sits alongside several older dated
copies in the same folder), 339 driver records, 98 raw McLeod columns.

Built specifically to give parse_samsara_drivers.py a real join key: Samsara's
scheduled reports are keyed by driver full name ("JERRY C. PARKS, JR"), not
by the driver code or tractor # the rest of this model uses -- this roster is
what lets that be a real name -> code lookup instead of a guess.
"""
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl

SRC = Path(__file__).parent / "data" / "raw" / "Driver_Data_latest.xlsx"
OUT = Path(__file__).parent / "output" / "driver_roster_data.json"

# These driver codes are status placeholders, not real people (matches the
# same set of pseudo-codes seen elsewhere in this model, e.g. parse_tractor_status.py's
# SOP_CODES) -- excluded from the roster so they never accidentally "match" a Samsara name.
NON_DRIVER_CODES = {
    "ASSIGNED", "AVAILABL", "PENDING", "PREP", "ORIENTAT", "TOBELP", "TOSELL",
    "WRECK", "SPARE", "MISC", "BOX", "SHOPTRUK", "OSHAD", "OSHAE", "UNKNOWN",
    "TDRIVER", "RDW", "RDWHOU", "RDWL", "RDWMAU", "RDWMOB", "RDWMTP", "RDWNMOB",
    "RDWRO", "RDWTCI", "RDWWIL",
}

SUFFIX_RE = re.compile(r"\s*,?\s*\b(JR|SR|II|III|IV)\.?\s*$", re.IGNORECASE)


def normalize_name(first, middle, last):
    """Build a matchable key: uppercase, no punctuation, no generational suffix,
    'FIRST LAST' (middle initial dropped -- Samsara names don't reliably carry one)."""
    last = SUFFIX_RE.sub("", str(last or "")).strip()
    parts = [str(first or "").strip(), last]
    key = " ".join(p for p in parts if p)
    key = re.sub(r"[^\w\s]", "", key).upper()
    key = re.sub(r"\s+", " ", key).strip()
    return key


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    ws = wb["DRIVER LISTING"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header) if h}

    def cell(r, col):
        i = idx.get(col)
        return r[i] if i is not None and i < len(r) else None

    drivers = {}
    name_index = {}
    for r in rows[1:]:
        code = str(cell(r, "Driver Code") or "").strip().upper()
        if not code or code in NON_DRIVER_CODES:
            continue
        first = str(cell(r, "First Name") or "").strip()
        last = str(cell(r, "Last Name") or "").strip()
        if not first and not last:
            continue
        rec = {
            "code": code,
            "firstName": first,
            "middleInitial": str(cell(r, "Mid.Initial") or "").strip(),
            "lastName": last,
            "fleet": str(cell(r, "Fleet") or "").strip(),
            "active": str(cell(r, "Active") or "").strip(),
            "status": str(cell(r, "Status") or "").strip(),
            "driverManager": str(cell(r, "Driver Manager") or "").strip(),
        }
        drivers[code] = rec
        key = normalize_name(first, rec["middleInitial"], last)
        if key:
            name_index.setdefault(key, []).append(code)

    ambiguous = {k: v for k, v in name_index.items() if len(v) > 1}

    out = {
        "checkedAt": datetime.now().isoformat(timespec="seconds"),
        "source": "McLeod Driver Listing export (Driver Data <date>.xlsx, SharePoint Sync)",
        "driverCount": len(drivers),
        "drivers": drivers,
        "nameIndex": {k: v[0] for k, v in name_index.items() if len(v) == 1},
        "ambiguousNameCount": len(ambiguous),
    }

    print(f"=== Driver Roster: {len(drivers)} drivers, {len(name_index)} unique name keys "
          f"({len(ambiguous)} ambiguous -- same normalized name, different codes) ===")

    OUT.write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(f"Wrote -> {OUT}")


if __name__ == "__main__":
    main()
